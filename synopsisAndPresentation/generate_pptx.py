import collections
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

BASE = 'c:/Users/nsraw/StudioProjects/dhfif/synopsisAndPresentation'

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    BG = RGBColor(0x0F, 0x0F, 0x14)
    CARD = RGBColor(0x1A, 0x1C, 0x2E)
    CARD2 = RGBColor(0x22, 0x24, 0x38)
    GREEN = RGBColor(0x00, 0xE6, 0x76)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    GRAY = RGBColor(0xAA, 0xAA, 0xBB)
    BLUE = RGBColor(0x44, 0xAA, 0xFF)
    RED = RGBColor(0xFF, 0x55, 0x55)
    ORANGE = RGBColor(0xFF, 0xAA, 0x33)
    PURPLE = RGBColor(0xBB, 0x77, 0xFF)
    DARK_GREEN = RGBColor(0x00, 0x44, 0x22)

    def set_bg(slide):
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = BG

    def rect(slide, l, t, w, h, color):
        s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
        s.fill.solid()
        s.fill.fore_color.rgb = color
        s.line.fill.background()
        return s

    def flat_rect(slide, l, t, w, h, color):
        s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
        s.fill.solid()
        s.fill.fore_color.rgb = color
        s.line.fill.background()
        return s

    def txt(slide, l, t, w, h, text, sz=18, color=WHITE, bold=False, align=PP_ALIGN.LEFT):
        tb = slide.shapes.add_textbox(l, t, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(sz)
        p.font.color.rgb = color
        p.font.bold = bold
        p.font.name = 'Calibri'
        p.alignment = align
        return tb

    def badge_txt(slide, l, t, w, h, text, sz, bg_color, fg_color=None):
        b = rect(slide, l, t, w, h, bg_color)
        tf = b.text_frame
        tf.paragraphs[0].text = text
        tf.paragraphs[0].font.size = Pt(sz)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = fg_color or BG
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        return b

    def title_bar(slide, title, subtitle=None):
        flat_rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(1.2), CARD)
        line = flat_rect(slide, Inches(0), Inches(1.16), Inches(13.333), Inches(0.05), GREEN)
        txt(slide, Inches(0.5), Inches(0.15), Inches(12), Inches(0.65), title, 30, GREEN, True)
        if subtitle:
            txt(slide, Inches(0.5), Inches(0.7), Inches(12), Inches(0.4), subtitle, 15, GRAY)

    def dot(slide, l, t, color=GREEN):
        d = slide.shapes.add_shape(MSO_SHAPE.OVAL, l, t, Inches(0.15), Inches(0.15))
        d.fill.solid()
        d.fill.fore_color.rgb = color
        d.line.fill.background()

    def bullets(slide, items, top=Inches(1.6), sz=16, gap=0.6):
        for i, item in enumerate(items):
            y = top + Inches(i * gap)
            rect(slide, Inches(0.6), y, Inches(12.1), Inches(gap - 0.08), CARD)
            dot(slide, Inches(0.85), y + Inches((gap - 0.08) / 2 - 0.07))
            txt(slide, Inches(1.15), y + Inches(0.06), Inches(11.3), Inches(gap - 0.15), item, sz, WHITE)

    # ===================================================================
    # SLIDE 1: TITLE
    # ===================================================================
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    rect(s, Inches(1.2), Inches(0.6), Inches(10.9), Inches(6.2), CARD)
    rect(s, Inches(1.6), Inches(1.0), Inches(10.1), Inches(1.8), DARK_GREEN)
    txt(s, Inches(1.8), Inches(1.1), Inches(9.7), Inches(0.9),
        "FitFi: AI-Powered Smart Fitness Coach", 38, GREEN, True, PP_ALIGN.CENTER)
    txt(s, Inches(1.8), Inches(1.95), Inches(9.7), Inches(0.5),
        "A Comprehensive Mobile Fitness and Wellness Application", 18, WHITE, False, PP_ALIGN.CENTER)
    txt(s, Inches(1.8), Inches(2.5), Inches(9.7), Inches(0.4),
        "MAJOR PROJECT - I  |  SYNOPSIS PRESENTATION  (REVIEW - I)", 15, BLUE, True, PP_ALIGN.CENTER)

    rect(s, Inches(1.8), Inches(3.3), Inches(4.8), Inches(2.3), CARD2)
    txt(s, Inches(2.0), Inches(3.4), Inches(4.4), Inches(0.3), "Submitted By:", 13, GREEN, True)
    tb = s.shapes.add_textbox(Inches(2.0), Inches(3.8), Inches(4.4), Inches(1.6))
    tf = tb.text_frame; tf.word_wrap = True
    for name in ["Student Name 1 (Roll No.)", "Student Name 2 (Roll No.)", "Student Name 3 (Roll No.)"]:
        p = tf.add_paragraph(); p.text = name; p.font.size = Pt(15); p.font.color.rgb = GRAY; p.font.name = 'Calibri'

    rect(s, Inches(7.0), Inches(3.3), Inches(4.8), Inches(2.3), CARD2)
    txt(s, Inches(7.2), Inches(3.4), Inches(4.4), Inches(0.3), "Supervisor:", 13, GREEN, True)
    tb2 = s.shapes.add_textbox(Inches(7.2), Inches(3.8), Inches(4.4), Inches(1.6))
    tf2 = tb2.text_frame; tf2.word_wrap = True
    for line in ["Prof. _______ (CSE)", "SoEC, DBUU, Dehradun"]:
        p = tf2.add_paragraph(); p.text = line; p.font.size = Pt(15); p.font.color.rgb = GRAY; p.font.name = 'Calibri'

    txt(s, Inches(1.2), Inches(6.15), Inches(10.9), Inches(0.4),
        "Dev Bhoomi Uttarakhand University  |  B.Tech CSE  |  7th Semester  |  Batch 2026-27",
        12, GRAY, False, PP_ALIGN.CENTER)

    # ===================================================================
    # SLIDE 2: TABLE OF CONTENTS
    # ===================================================================
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    title_bar(s, "Table of Contents")

    toc = [
        "Introduction", "Problem Statement", "Existing Systems & Limitations",
        "Development Gap", "Project Objectives", "System Architecture",
        "Module: AI Exercise Detection (Real-Time Camera AI)",
        "Module: Heart Rate Monitor & Step Counter",
        "Module: Water Tracker & Gym Workout System",
        "Module: UFC Combat Training & Challenge + XP System",
        "Technologies & Tools Used",
        "App Demo — Live Videos",
        "Work Plan & Timeline",
        "Outcome & Conclusion", "References",
    ]
    for i, item in enumerate(toc):
        top = Inches(1.4) + Inches(i * 0.39)
        badge_txt(s, Inches(0.7), top, Inches(0.4), Inches(0.3), str(i+1), 12, GREEN)
        txt(s, Inches(1.25), top + Inches(0.01), Inches(11), Inches(0.28), item, 16, WHITE)

    # ===================================================================
    # SLIDE 3: INTRODUCTION
    # ===================================================================
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    title_bar(s, "Introduction", "What is FitFi?")

    bullets(s, [
        "FitFi is a comprehensive, all-in-one fitness application built with the Flutter framework.",
        "It uses REAL on-device AI (Google ML Kit Pose Detection) to track exercise form via live camera.",
        "Heart rate is measured using the phone camera (PPG technique) — no smartwatch required.",
        "Steps are counted using the phone's built-in accelerometer sensor — real hardware sensing.",
        "Includes a library of 256 exercises — Gym courses, UFC combat training, and Challenges.",
        "Features a gamification system — earn XP, level up, unlock ranks (Solo Leveling inspired).",
        "Water tracker with automated reminders — notifications every 2 hours to stay hydrated.",
        "The entire app works 100% OFFLINE — all AI processing happens on-device.",
        "Supports multiple languages (English & Hindi).",
    ], Inches(1.5), 15, 0.6)

    # ===================================================================
    # SLIDE 4: PROBLEM STATEMENT
    # ===================================================================
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    title_bar(s, "Problem Statement", "Why was FitFi needed?")

    problems = [
        ("Fragmented Apps", "Users need 3-4 separate apps for steps, water, heart rate, workouts & diet tracking."),
        ("No Real-Time AI Feedback", "Existing workout apps show video tutorials but cannot detect if your form is correct."),
        ("External Hardware Required", "Heart rate monitoring requires smartwatches — phone camera PPG is underutilized."),
        ("Low User Engagement", "Fitness apps lack gaming-level gamification to keep users motivated long-term."),
        ("Internet Dependency", "Many fitness apps require constant internet connectivity for core features."),
    ]
    for i, (title, desc) in enumerate(problems):
        top = Inches(1.6) + Inches(i * 1.1)
        rect(s, Inches(0.6), top, Inches(12.1), Inches(0.95), CARD)
        badge_txt(s, Inches(0.85), top + Inches(0.25), Inches(0.4), Inches(0.4), "X", 16, RED, WHITE)
        txt(s, Inches(1.45), top + Inches(0.08), Inches(10.8), Inches(0.35), title, 18, GREEN, True)
        txt(s, Inches(1.45), top + Inches(0.45), Inches(10.8), Inches(0.42), desc, 14, GRAY)

    # ===================================================================
    # SLIDE 5: EXISTING SYSTEMS
    # ===================================================================
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    title_bar(s, "Existing Systems & Their Limitations")

    apps = [
        ("Google Fit", "Step tracking & basic activity monitoring", "No AI form detection, no workout courses, no gamification", BLUE),
        ("MyFitnessPal", "Calorie tracking with large food database", "No exercise form detection, no camera heart rate, manual entry only", ORANGE),
        ("Nike Training Club", "Professional trainer-guided workout videos", "No real-time feedback on user form, cannot auto-count reps", RED),
        ("Freeletics", "AI-generated personalized workout plans", "AI only for plan generation, not real-time pose analysis during exercise", PURPLE),
    ]
    for i, (name, strength, weakness, color) in enumerate(apps):
        top = Inches(1.6) + Inches(i * 1.35)
        rect(s, Inches(0.6), top, Inches(12.1), Inches(1.2), CARD)
        badge_txt(s, Inches(0.85), top + Inches(0.15), Inches(2.2), Inches(0.35), name, 15, color, WHITE)
        txt(s, Inches(3.3), top + Inches(0.12), Inches(9.2), Inches(0.35), "Strength: " + strength, 14, GREEN)
        txt(s, Inches(3.3), top + Inches(0.55), Inches(9.2), Inches(0.55), "Limitation: " + weakness, 14, RED)

    txt(s, Inches(0.6), Inches(7.0), Inches(12.1), Inches(0.35),
        "No single app combines AI detection + hardware sensors + structured workouts + gamification.",
        15, ORANGE, True, PP_ALIGN.CENTER)

    # ===================================================================
    # SLIDE 6: DEVELOPMENT GAP
    # ===================================================================
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    title_bar(s, "Development Gap", "What is missing in the market that FitFi addresses?")

    gaps = [
        ("Fragmented Apps", "3-4 apps needed for different fitness activities", "FitFi integrates everything — steps, water, heart rate, AI detection, gym, UFC, diet."),
        ("No AI Form Check", "Pose detection tech exists but not in fitness apps", "Google ML Kit detects 33 body landmarks in real-time via camera."),
        ("No Phone Heart Rate", "PPG via phone camera is feasible but unused", "FitFi measures heart rate using rear camera + flashlight."),
        ("No Gamification", "Gaming-level engagement missing in fitness apps", "Solo Leveling inspired XP system — 10 levels, 11 ranks, celebrations."),
        ("Internet Required", "Most apps don't work offline", "FitFi operates 100% offline — all processing on-device."),
    ]
    for i, (gap, problem, solution) in enumerate(gaps):
        top = Inches(1.5) + Inches(i * 1.1)
        rect(s, Inches(0.6), top, Inches(12.1), Inches(0.95), CARD)
        txt(s, Inches(0.9), top + Inches(0.06), Inches(2.5), Inches(0.3), gap, 15, ORANGE, True)
        txt(s, Inches(0.9), top + Inches(0.35), Inches(5.3), Inches(0.25), "Gap: " + problem, 12, RED)
        txt(s, Inches(6.5), top + Inches(0.35), Inches(5.9), Inches(0.55), "FitFi: " + solution, 12, GREEN)

    # ===================================================================
    # SLIDE 7: OBJECTIVES
    # ===================================================================
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    title_bar(s, "Project Objectives")

    bullets(s, [
        "Develop a unified Flutter app combining steps, water, heart rate, AI detection, gym, UFC & diet.",
        "Implement real-time on-device AI pose detection (Google ML Kit) for rep counting & posture evaluation.",
        "Integrate camera-based PPG for heart rate measurement using phone's rear camera + flashlight.",
        "Implement accelerometer-based step counting with distance & calorie estimation.",
        "Build a library of 256 exercises with animated visual guides (Lottie, GIF, images).",
        "Create a gamified XP and ranking system (10 levels, 11 ranks) to boost user engagement.",
        "Provide structured gym courses, UFC combat training & a 6-level ranked challenge system.",
        "Ensure the entire application works offline — no internet dependency for core features.",
    ], Inches(1.5), 15, 0.68)

    # ===================================================================
    # SLIDE 8: ARCHITECTURE
    # ===================================================================
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    title_bar(s, "System Architecture", "Layered modular design for scalability")

    layers = [
        ("UI", "Presentation Layer", "Screens + Widgets + CustomPainter (arcs, waves, ECG, skeleton overlay)", GREEN),
        ("Logic", "Service Layer", "PoseDetection, RepCounter, StepCounter, HeartRate, Water, XP, Notification Services", BLUE),
        ("Data", "Data Layer", "256 exercises, gym courses, UFC programs, challenges, diet plans (Dart data files)", ORANGE),
        ("Store", "Storage Layer", "SharedPreferences — all data stored locally for offline operation", PURPLE),
        ("Perf", "Performance Layer", "Device profiling, adaptive memory/animation management, task concurrency", GRAY),
    ]
    for i, (tag, name, desc, color) in enumerate(layers):
        top = Inches(1.6) + Inches(i * 1.05)
        rect(s, Inches(0.6), top, Inches(12.1), Inches(0.9), CARD)
        flat_rect(s, Inches(0.6), top, Inches(0.07), Inches(0.9), color)
        badge_txt(s, Inches(0.85), top + Inches(0.25), Inches(0.7), Inches(0.35), tag, 11, color)
        txt(s, Inches(1.75), top + Inches(0.07), Inches(10.5), Inches(0.35), name, 18, color, True)
        txt(s, Inches(1.75), top + Inches(0.45), Inches(10.5), Inches(0.4), desc, 13, GRAY)

    for i in range(4):
        top = Inches(1.6) + Inches((i + 1) * 1.05) - Inches(0.15)
        ar = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(6.5), top, Inches(0.35), Inches(0.15))
        ar.fill.solid(); ar.fill.fore_color.rgb = GREEN; ar.line.fill.background()

    # ===================================================================
    # SLIDE 9: AI EXERCISE DETECTION
    # ===================================================================
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    title_bar(s, "Module: AI Exercise Detection", "The core AI feature — Real-time Computer Vision via Google ML Kit")

    rect(s, Inches(0.4), Inches(1.5), Inches(6.3), Inches(5.5), CARD)
    txt(s, Inches(0.6), Inches(1.6), Inches(5.9), Inches(0.35), "How It Works:", 20, GREEN, True)
    steps_ai = [
        "1. Camera opens — front or rear lens",
        "2. Google ML Kit Pose Detection identifies 33 body landmarks in REAL TIME",
        "3. Joint angles calculated using atan2 trigonometry",
        "4. Push-up: Elbow <90° = down, >150° = up → 1 rep counted",
        "5. Squat: Knee <100° = down, >158° = up → 1 rep counted",
        "6. GREEN skeleton overlay when posture is correct",
        "7. RED skeleton + feedback: \"Go lower\", \"Push up\", \"Fix Position\"",
    ]
    for i, step in enumerate(steps_ai):
        txt(s, Inches(0.7), Inches(2.1 + i * 0.6), Inches(5.8), Inches(0.5), step, 13, WHITE)

    rect(s, Inches(6.9), Inches(1.5), Inches(6.0), Inches(5.5), CARD)
    txt(s, Inches(7.1), Inches(1.6), Inches(5.6), Inches(0.35), "Supported Exercises:", 20, BLUE, True)

    txt(s, Inches(7.1), Inches(2.15), Inches(5.6), Inches(0.3), "Repetition-Based (Auto Counting):", 15, ORANGE, True)
    rep_exs = ["Push-ups", "Pull-ups", "Squats", "Lunges", "Crunches", "Burpees", "Jumping Jacks", "High Knees"]
    for i, ex in enumerate(rep_exs):
        col = i % 2; row = i // 2
        txt(s, Inches(7.2 + col * 2.7), Inches(2.55 + row * 0.4), Inches(2.5), Inches(0.35), "• " + ex, 13, WHITE)

    txt(s, Inches(7.1), Inches(4.25), Inches(5.6), Inches(0.3), "Isometric Hold-Based (Timer):", 15, PURPLE, True)
    hold_exs = ["Plank (alignment >150°)", "Wall Sit (knee 70°-115°)", "Side Plank", "V-Sit"]
    for i, ex in enumerate(hold_exs):
        col = i % 2; row = i // 2
        txt(s, Inches(7.2 + col * 2.7), Inches(4.65 + row * 0.4), Inches(2.5), Inches(0.35), "• " + ex, 13, WHITE)

    txt(s, Inches(7.1), Inches(5.7), Inches(5.6), Inches(0.4),
        "Also used in Challenge mode for AI-validated completion.", 13, GREEN, True)

    # ===================================================================
    # SLIDE 10: HEART RATE + STEP COUNTER
    # ===================================================================
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    title_bar(s, "Module: Heart Rate Monitor & Step Counter", "Real hardware sensor integration")

    rect(s, Inches(0.4), Inches(1.5), Inches(6.3), Inches(5.5), CARD)
    txt(s, Inches(0.6), Inches(1.6), Inches(5.9), Inches(0.35), "Heart Rate Monitor (Camera PPG)", 20, RED, True)
    hr = [
        "• Uses rear camera + flashlight (torch mode)",
        "• User places fingertip on camera lens",
        "• Blood volume changes cause light intensity variations",
        "• Luminance analyzed in 50×50 pixel ROI per frame",
        "• Moving average threshold detects pulse peaks",
        "• Real-time BPM + signal quality calculated",
        "",
        "  3-Step UI Flow:",
        "  1. Instruction — \"Place finger on flashlight\"",
        "  2. Scanning — Animated radar + 10s countdown",
        "  3. Result — BPM, stress level, recovery, weekly trends",
    ]
    for i, line in enumerate(hr):
        if line == "": continue
        txt(s, Inches(0.7), Inches(2.1 + i * 0.45), Inches(5.7), Inches(0.4), line, 13,
            GRAY if line.startswith("  ") else WHITE)

    rect(s, Inches(6.9), Inches(1.5), Inches(6.0), Inches(5.5), CARD)
    txt(s, Inches(7.1), Inches(1.6), Inches(5.6), Inches(0.35), "Step Counter (Accelerometer)", 20, GREEN, True)
    sc = [
        "• Uses device accelerometer via sensors_plus package",
        "• Streams 3D acceleration data (X, Y, Z axes)",
        "• Vector magnitude threshold > 12.0 m/s² = 1 step",
        "• 300ms debounce filter prevents false positives",
        "• Estimates distance (km) and calories burned",
        "",
        "  Screen Features:",
        "  • Animated circular arc progress (10,000 step goal)",
        "  • Weekly / Monthly averages with toggle tabs",
        "  • Trend chart with cubic Bezier curves",
        "  • Recent history with badges (Goal Reached, Active)",
    ]
    for i, line in enumerate(sc):
        if line == "": continue
        txt(s, Inches(7.2), Inches(2.1 + i * 0.45), Inches(5.5), Inches(0.4), line, 13,
            GRAY if line.startswith("  ") else WHITE)

    # ===================================================================
    # SLIDE 11: WATER + GYM
    # ===================================================================
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    title_bar(s, "Module: Water Tracker & Gym Workout System")

    rect(s, Inches(0.4), Inches(1.5), Inches(6.3), Inches(5.5), CARD)
    txt(s, Inches(0.6), Inches(1.6), Inches(5.9), Inches(0.35), "Water Tracker", 20, BLUE, True)
    wt = [
        "3 vessel sizes: Cup (150ml), Glass (250ml), Bottle (500ml)",
        "Animated sinusoidal wave fill painter for liquid level",
        "Customizable daily water intake goal",
        "Weekly average bar chart + monthly completion ring",
        "Automated reminders every 2 hours (8 AM - 8 PM)",
        "Timestamped intake history log",
        "One-tap logging with scale pulse animation",
    ]
    for i, line in enumerate(wt):
        dot(s, Inches(0.75), Inches(2.15 + i * 0.6), BLUE)
        txt(s, Inches(1.05), Inches(2.05 + i * 0.6), Inches(5.5), Inches(0.5), line, 14, WHITE)

    rect(s, Inches(6.9), Inches(1.5), Inches(6.0), Inches(5.5), CARD)
    txt(s, Inches(7.1), Inches(1.6), Inches(5.6), Inches(0.35), "Gym Workout System", 20, ORANGE, True)
    gym = [
        "8-step personalization onboarding wizard",
        "256 exercises with Lottie / GIF / image animations",
        "5-phase guided workout player with timers",
        "Custom workout builder from exercise library",
        "Progressive overload algorithm (reps +30%, rest -10s)",
        "Body-type & focus area course filtering",
        "Victory dialog and XP rewards on completion",
    ]
    for i, line in enumerate(gym):
        dot(s, Inches(7.25), Inches(2.15 + i * 0.6), ORANGE)
        txt(s, Inches(7.55), Inches(2.05 + i * 0.6), Inches(5.2), Inches(0.5), line, 14, WHITE)

    # ===================================================================
    # SLIDE 12: UFC + CHALLENGE
    # ===================================================================
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    title_bar(s, "Module: UFC Combat Training & Challenge + XP System")

    rect(s, Inches(0.4), Inches(1.5), Inches(6.3), Inches(5.5), CARD)
    txt(s, Inches(0.6), Inches(1.6), Inches(5.9), Inches(0.35), "UFC Combat Training", 20, RED, True)
    ufc_items = [
        ("Two distinct training styles:", WHITE, False),
        ("  Dagestani (Khabib): Wrestling, grappling, cardio", GRAY, False),
        ("  Irish (McGregor): Striking, kickboxing, power", GRAY, False),
        ("Multiple courses per training style", WHITE, False),
        ("6 training days + 1 rest day per program", WHITE, False),
        ("Exercises with animated thumbnails & sets × reps", WHITE, False),
        ("Launches guided WorkoutPlayer for each session", WHITE, False),
        ("Completion tracked via ProgressService", WHITE, False),
    ]
    for i, (line, color, _) in enumerate(ufc_items):
        if not line.startswith("  "):
            dot(s, Inches(0.75), Inches(2.15 + i * 0.55), RED)
        txt(s, Inches(1.05 if not line.startswith("  ") else 0.9), Inches(2.05 + i * 0.55),
            Inches(5.5), Inches(0.45), line, 13, color)

    rect(s, Inches(6.9), Inches(1.5), Inches(6.0), Inches(5.5), CARD)
    txt(s, Inches(7.1), Inches(1.6), Inches(5.6), Inches(0.35), "Challenge System + XP Gamification", 20, GREEN, True)
    txt(s, Inches(7.1), Inches(2.1), Inches(5.6), Inches(0.3), "6 Challenge Levels:", 15, ORANGE, True)
    levels = ["Lvl 10: Junior Trainee", "Lvl 20: Senior Trainee", "Lvl 30: Gym Bro",
              "Lvl 50: Beast Mode", "Lvl 60: Apex Predator", "Lvl 100: UFC Fighter"]
    for i, lv in enumerate(levels):
        col = i % 2; row = i // 2
        txt(s, Inches(7.2 + col * 2.7), Inches(2.5 + row * 0.35), Inches(2.5), Inches(0.3), lv, 12, WHITE)

    txt(s, Inches(7.1), Inches(3.65), Inches(5.6), Inches(0.3),
        "15 individual quick challenges — validated via AI camera", 13, BLUE)

    txt(s, Inches(7.1), Inches(4.1), Inches(5.6), Inches(0.3), "XP System (Solo Leveling Inspired):", 15, PURPLE, True)
    xp_items = [
        "10 XP levels (100 to 100,000 XP thresholds)",
        "11 ranks: Unranked → F → E → ... → SSS → X Rank",
        "Dual domain tracks: Gym XP + Challenge XP",
        "Level-up celebration dialogs with rank badges",
    ]
    for i, line in enumerate(xp_items):
        dot(s, Inches(7.25), Inches(4.55 + i * 0.45), PURPLE)
        txt(s, Inches(7.55), Inches(4.45 + i * 0.45), Inches(5.2), Inches(0.4), line, 13, WHITE)

    # ===================================================================
    # SLIDE 13: TECHNOLOGIES
    # ===================================================================
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    title_bar(s, "Technologies & Tools Used")

    tech = [
        ("Language", "Dart", "Primary language for Flutter development", GREEN),
        ("Framework", "Flutter", "Cross-platform mobile app (Android + iOS)", GREEN),
        ("AI / ML", "Google ML Kit Pose Detection", "On-device 33-landmark body tracking in real-time", BLUE),
        ("Camera", "camera (Flutter package)", "Live camera for AI pose detection + PPG heart rate", BLUE),
        ("Sensors", "sensors_plus", "Accelerometer data stream for step counting", BLUE),
        ("Storage", "SharedPreferences", "Local offline data persistence", ORANGE),
        ("Notifications", "flutter_local_notifications", "Scheduled water & workout reminders", ORANGE),
        ("Animations", "Lottie", "Rich exercise animations (After Effects JSON)", PURPLE),
        ("UI", "percent_indicator, google_fonts", "Progress rings + custom typography", GRAY),
        ("Calendar", "table_calendar", "Interactive workout history calendar view", GRAY),
        ("IDE / VCS", "Android Studio + Git", "Development environment + version control", GRAY),
    ]
    for i, (cat, tool, desc, color) in enumerate(tech):
        top = Inches(1.45) + Inches(i * 0.51)
        rect(s, Inches(0.5), top, Inches(12.3), Inches(0.43), CARD)
        badge_txt(s, Inches(0.65), top + Inches(0.06), Inches(1.6), Inches(0.3), cat, 11, color)
        txt(s, Inches(2.45), top + Inches(0.06), Inches(3.5), Inches(0.3), tool, 14, WHITE, True)
        txt(s, Inches(6.0), top + Inches(0.06), Inches(6.5), Inches(0.3), desc, 12, GRAY)

    # ===================================================================
    # SLIDE 14: APP DEMO VIDEOS
    # ===================================================================
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    title_bar(s, "App Demo — Live Videos", "Click on the video to play during presentation")

    poster1 = os.path.join(BASE, 'poster1.png')
    poster2 = os.path.join(BASE, 'poster2.png')
    vid1 = os.path.join(BASE, 'fitfi app', 'WhatsApp Video 2026-09-13 at 2.36.10 PM.mp4')
    vid2 = os.path.join(BASE, 'fitfi app', 'ffr.mp4')

    # Video 1
    rect(s, Inches(0.4), Inches(1.5), Inches(6.2), Inches(5.5), CARD)
    txt(s, Inches(0.6), Inches(1.6), Inches(5.8), Inches(0.35), "Video 1: App Demo", 20, GREEN, True)
    try:
        s.shapes.add_movie(vid1, Inches(0.7), Inches(2.1), Inches(5.6), Inches(4.5),
                           poster_frame_image=poster1, mime_type='video/mp4')
    except Exception as e:
        print(f"Video 1 embed note: {e}")
        rect(s, Inches(0.7), Inches(2.1), Inches(5.6), Inches(4.5), CARD2)
        txt(s, Inches(1.5), Inches(4.0), Inches(4.0), Inches(0.5),
            "Video: WhatsApp Video 2026-09-13.mp4", 14, GRAY, False, PP_ALIGN.CENTER)

    # Video 2
    rect(s, Inches(6.8), Inches(1.5), Inches(6.2), Inches(5.5), CARD)
    txt(s, Inches(7.0), Inches(1.6), Inches(5.8), Inches(0.35), "Video 2: Full Feature Recording", 20, BLUE, True)
    try:
        s.shapes.add_movie(vid2, Inches(7.1), Inches(2.1), Inches(5.6), Inches(4.5),
                           poster_frame_image=poster2, mime_type='video/mp4')
    except Exception as e:
        print(f"Video 2 embed note: {e}")
        rect(s, Inches(7.1), Inches(2.1), Inches(5.6), Inches(4.5), CARD2)
        txt(s, Inches(8.0), Inches(4.0), Inches(4.0), Inches(0.5),
            "Video: ffr.mp4", 14, GRAY, False, PP_ALIGN.CENTER)

    # ===================================================================
    # SLIDE 15: TIMELINE
    # ===================================================================
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    title_bar(s, "Work Plan & Timeline")

    months = ["Month 1", "Month 2", "Month 3", "Month 4"]
    for i, m in enumerate(months):
        badge_txt(s, Inches(6.5 + i * 1.6), Inches(1.5), Inches(1.4), Inches(0.35), m, 12, GREEN)

    tasks = [
        ("Requirement Analysis + UI Design", [1,0,0,0]),
        ("Project Setup + Architecture", [1,0,0,0]),
        ("Step Counter + Water Tracker + Heart Rate", [0,1,0,0]),
        ("AI Pose Detection + Rep Counter", [0,1,1,0]),
        ("Gym Courses + Workout Player", [0,0,1,0]),
        ("UFC Training + Challenges", [0,0,1,0]),
        ("Diet Plan + Account Dashboard", [0,0,1,0]),
        ("XP / Gamification System", [0,0,1,0]),
        ("Performance Optimization + Testing", [0,0,0,1]),
        ("Documentation + Presentation", [0,0,0,1]),
    ]
    for i, (task, active) in enumerate(tasks):
        top = Inches(2.05) + Inches(i * 0.5)
        txt(s, Inches(0.5), top + Inches(0.03), Inches(5.8), Inches(0.35), task, 13, WHITE)
        for j, a in enumerate(active):
            left = Inches(6.5 + j * 1.6)
            if a:
                rect(s, left + Inches(0.1), top + Inches(0.07), Inches(1.2), Inches(0.28), GREEN)
            else:
                rect(s, left + Inches(0.1), top + Inches(0.07), Inches(1.2), Inches(0.28), CARD)

    txt(s, Inches(0.5), Inches(7.1), Inches(12.3), Inches(0.3),
        "All modules completed — Documentation and final presentation in progress.",
        13, GREEN, True, PP_ALIGN.CENTER)

    # ===================================================================
    # SLIDE 16: OUTCOME & CONCLUSION
    # ===================================================================
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    title_bar(s, "Outcome & Conclusion")

    bullets(s, [
        "Built a fully functional fitness app with 10+ health modules integrated in one unified platform.",
        "Real on-device AI — live camera pose detection counts reps & evaluates form WITHOUT internet.",
        "Camera-based PPG heart rate + accelerometer step counter — no external wearable hardware needed.",
        "256 exercises with animated guides across Gym, UFC, Home Workout & Challenge categories.",
        "Gamified XP / ranking system with 10 levels & 11 ranks keeps users motivated and consistent.",
        "Entire application works 100% offline — all AI and data processing happens on-device.",
        "Future-ready: Python ML models (body type, diet, planner) already trained — ready for TFLite integration.",
        "Modular architecture enables seamless addition of cloud sync, social features & wearable support.",
    ], Inches(1.5), 15, 0.68)

    # ===================================================================
    # SLIDE 17: REFERENCES
    # ===================================================================
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    title_bar(s, "References")

    refs = [
        '[1] Google — "ML Kit Pose Detection" — developers.google.com/ml-kit/vision/pose-detection',
        '[2] Google — "Flutter Framework" — flutter.dev',
        '[3] "sensors_plus: Accelerometer plugin" — pub.dev/packages/sensors_plus',
        '[4] "camera: Flutter camera plugin" — pub.dev/packages/camera',
        '[5] "google_mlkit_pose_detection" — pub.dev/packages/google_mlkit_pose_detection',
        '[6] J. Allen — "Photoplethysmography in clinical measurement" — Physiological Measurement, 2007',
        '[7] "Lottie for Flutter" — pub.dev/packages/lottie',
        '[8] "flutter_local_notifications" — pub.dev/packages/flutter_local_notifications',
        '[9] "percent_indicator" — pub.dev/packages/percent_indicator',
        '[10] "table_calendar" — pub.dev/packages/table_calendar',
    ]
    for i, ref in enumerate(refs):
        txt(s, Inches(0.8), Inches(1.5 + i * 0.55), Inches(11.7), Inches(0.45), ref, 14, GRAY)

    # ===================================================================
    # SLIDE 18: THANK YOU
    # ===================================================================
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s)
    rect(s, Inches(2.0), Inches(1.2), Inches(9.3), Inches(5.0), CARD)
    rect(s, Inches(2.5), Inches(1.8), Inches(8.3), Inches(1.5), DARK_GREEN)
    txt(s, Inches(2.5), Inches(2.0), Inches(8.3), Inches(1.0),
        "Thank You!", 52, GREEN, True, PP_ALIGN.CENTER)
    txt(s, Inches(2.5), Inches(3.6), Inches(8.3), Inches(0.6),
        "Any Questions?", 28, WHITE, False, PP_ALIGN.CENTER)
    txt(s, Inches(2.5), Inches(4.5), Inches(8.3), Inches(0.5),
        "FitFi: AI-Powered Smart Fitness Coach", 16, GRAY, False, PP_ALIGN.CENTER)
    txt(s, Inches(2.5), Inches(5.2), Inches(8.3), Inches(0.4),
        "Dev Bhoomi Uttarakhand University  |  B.Tech CSE  |  7th Semester", 13, GRAY, False, PP_ALIGN.CENTER)

    # ===== SAVE =====
    out_path = os.path.join(BASE, 'FitFi_Presentation_v3.pptx')
    prs.save(out_path)
    print(f"Presentation generated: {len(prs.slides)} slides → {out_path}")

if __name__ == '__main__':
    create_presentation()
