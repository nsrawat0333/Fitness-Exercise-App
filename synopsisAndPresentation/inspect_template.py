import pptx

prs = pptx.Presentation('c:/Users/nsraw/StudioProjects/dhfif/synopsisAndPresentation/PPT Format.pptx')
for idx, slide in enumerate(prs.slides):
    print(f"=== Slide {idx+1} (layout: {slide.slide_layout.name}) ===")
    for sh in slide.shapes:
        t = sh.text[:40].replace('\n', ' ') if sh.has_text_frame else ''
        print(f"  Shape '{sh.name}' ({sh.shape_type}): {t}")
